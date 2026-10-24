from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path
# compact_reference_guide; memo_masthead. Override: monochrome headings, compact 10pt tables.
d=Document(); s=d.sections[0]
s.page_width=Inches(8.5); s.page_height=Inches(11)
s.top_margin=s.bottom_margin=s.left_margin=s.right_margin=Inches(1)
s.header_distance=s.footer_distance=Inches(.492)
normal=d.styles['Normal']; normal.font.name='Calibri'; normal.font.size=Pt(11)
normal.paragraph_format.space_before=Pt(0); normal.paragraph_format.space_after=Pt(6); normal.paragraph_format.line_spacing=1.25
for name,size,before,after in [('Title',18,0,8),('Heading 1',16,18,10),('Heading 2',13,14,7)]:
 st=d.styles[name]; st.font.name='Calibri'; st.font.size=Pt(size); st.font.color.rgb=RGBColor(0,0,0); st.font.bold=True
 st.paragraph_format.space_before=Pt(before); st.paragraph_format.space_after=Pt(after)
s.header.paragraphs[0].text='SAO ĐỎ GROUP | PHỤ LỤC TÀI SẢN IT'
s.header.paragraphs[0].runs[0].font.size=Pt(9)
f=s.footer.paragraphs[0]; f.text='Dữ liệu tham khảo • Trang '
f._p.append(OxmlElement('w:fldSimple')); f._p[-1].set(qn('w:instr'),'PAGE')
def table(rows,widths):
 t=d.add_table(rows=0,cols=len(widths)); t.autofit=False
 pr=t._tbl.tblPr; w=pr.find(qn('w:tblW')); w.set(qn('w:w'),'9360'); w.set(qn('w:type'),'dxa')
 ind=OxmlElement('w:tblInd'); ind.set(qn('w:w'),'120'); ind.set(qn('w:type'),'dxa'); pr.append(ind)
 for col,width in zip(t.columns,widths): col.width=Inches(width/1440)
 for i,row in enumerate(rows):
  cells=t.add_row().cells; trpr=t.rows[-1]._tr.get_or_add_trPr(); trpr.append(OxmlElement('w:cantSplit'))
  if i==0: trpr.append(OxmlElement('w:tblHeader'))
  for c,width,txt in zip(cells,widths,row):
   c.width=Inches(width/1440); c.text=txt; cp=c._tc.get_or_add_tcPr()
   margins=OxmlElement('w:tcMar')
   for edge,value in [('top',80),('bottom',80),('start',120),('end',120)]:
    e=OxmlElement('w:'+edge); e.set(qn('w:w'),str(value)); e.set(qn('w:type'),'dxa'); margins.append(e)
   cp.append(margins)
   borders=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    e=OxmlElement('w:'+edge); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'AAAAAA'); borders.append(e)
   cp.append(borders)
   if i==0:
    sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'E8EEF5'); cp.append(sh)
   for p in c.paragraphs:
    p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1
    for r in p.runs: r.font.size=Pt(10); r.bold=i==0
 return t

d.add_paragraph('DANH SÁCH TÀI SẢN IT\nVÀ GHI CHÚ BỔ SUNG', 'Title')
d.add_paragraph('Tài liệu đính kèm • Tổng hợp từ ảnh biên bản BMHC 0501 do người phụ trách cung cấp.')
d.add_paragraph('Người giao ghi trên ảnh: Nguyễn Minh Dương – IT. Ngày bàn giao và bên nhận chưa được điền. Số liệu dưới đây dùng để tham khảo, chưa xác nhận qua kiểm kê từng thiết bị.')
items=[
('Máy tính lắp ráp',33,'Bổ sung cấu hình, người dùng và vị trí từng máy.'),
('Laptop',19,'Bổ sung hãng/model, serial và người sử dụng.'),
('Bộ phát Wifi',10,'Bổ sung model và khu vực lắp đặt.'),
('Tổng đài',1,'Bổ sung model, vị trí và tình trạng hoạt động.'),
('Điện thoại bàn',8,'Bổ sung phòng sử dụng và số máy nhánh nếu có.'),
('Máy chủ Dell T360',1,'Đối chiếu model, service tag và vị trí đặt máy.'),
('UPS online MHPRO 3000',1,'Đối chiếu model, vị trí và tình trạng ắc quy.'),
('Cân bằng tải Rujie NBR6205-E',1,'Tên chép theo ảnh; cần đối chiếu hãng/model trên nhãn.'),
('Switch cisco 24 Port 100/1000',1,'Tên chép theo ảnh; bổ sung model và vị trí lắp đặt.'),
('Máy in',12,'Bổ sung hãng/model, phòng sử dụng và tình trạng.'),
('Máy Scan',3,'Bổ sung hãng/model, vị trí và người/phòng sử dụng.'),
]
assert sum(x[1] for x in items)==90
rows=[('STT','Thiết bị theo ảnh','SL','Ghi chú cần bổ sung')]+[(str(i),name,str(n),note) for i,(name,n,note) in enumerate(items,1)]+[('','TỔNG','90','11 nhóm thiết bị')]
table(rows,[550,3100,500,5210])
p=d.add_paragraph('Ghi chú là đề xuất thông tin cần thu thập, không phải mô tả tình trạng thiết bị đã được xác nhận.'); p.paragraph_format.space_before=Pt(4); p.paragraph_format.space_after=Pt(4)
for r in p.runs:r.font.size=Pt(9)
d.add_page_break()
d.add_paragraph('Mẫu bổ sung từng tài sản', 'Title')
d.add_paragraph('Dùng một phiếu cho mỗi thiết bị thực tế khi kiểm kê. Chưa tạo sẵn 90 hồ sơ vì ảnh chỉ có số lượng theo nhóm. Trường chưa xác định được có thể bổ sung sau.')
table([('Thông tin','Nội dung ghi nhận'),('Mã tài sản nội bộ','................................................................'),('Tên thiết bị / nhóm ở trang 1','................................................................'),('Hãng, model / serial nếu có','................................................................'),('Phòng quản lý / người sử dụng','................................................................'),('Vị trí hiện tại','................................................................'),('Tình trạng quan sát được','................................................................'),('Vấn đề cần sửa hoặc cần báo','................................................................\n................................................................'),('Lịch sử sửa chữa đã biết','................................................................\n................................................................'),('Ghi chú khác','................................................................\n................................................................'),('Người ghi nhận / ngày ghi nhận','................................................................')],[2700,6660])
d.add_paragraph('Nguyên tắc ghi nhận', 'Heading 2')
d.add_paragraph('Ưu tiên đủ số lượng, xác định thiết bị và ghi nhận nhu cầu sửa chữa. Không tự suy người giữ, vị trí, tình trạng hoặc lịch sử bảo trì từ ảnh bàn giao. Nếu số lượng thực tế khác bảng tham khảo, ghi rõ kết quả kiểm kê và ngày ghi nhận.')
p=Path('docs/bao-cao/DANH-SACH-TAI-SAN-IT-VA-GHI-CHU.docx'); d.save(p)
check=Document(p); assert len(check.tables)==2; assert len(check.tables[0].rows)==13
print(p)
