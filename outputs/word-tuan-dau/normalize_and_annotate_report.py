from pathlib import Path
from shutil import copy2
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt
from lxml import etree


SOURCE = Path("docs/bao-cao/BAO-CAO-TUAN-DAU-DHN.docx")
OUTPUT = SOURCE
BACKUP = Path("outputs/word-tuan-dau/backups/BAO-CAO-TUAN-DAU-DHN-TRUOC-DONG-BO-FONT-VA-GHI-CHU.docx")
FONT = "Times New Roman"
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def package_text_signature(path: Path):
    signature = {}
    with ZipFile(path) as package:
        names = package.namelist()
        selected = [
            name
            for name in names
            if name == "word/document.xml"
            or name.startswith("word/header")
            or name.startswith("word/footer")
            or name in {"word/comments.xml", "word/footnotes.xml", "word/endnotes.xml"}
        ]
        for name in sorted(selected):
            root = etree.fromstring(package.read(name))
            signature[name] = tuple(root.xpath("//w:t/text()", namespaces={"w": W_NS}))
        media = {
            name: package.getinfo(name).file_size
            for name in names
            if name.startswith("word/media/")
        }
    return signature, media


def is_subsequence(original, reviewed):
    cursor = iter(reviewed)
    return all(any(candidate == item for candidate in cursor) for item in original)


def set_run_font(run, size):
    run.font.name = FONT
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = etree.SubElement(rpr, qn("w:rFonts"))
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{key}"), FONT)
    run.font.size = Pt(size)


def append_note(paragraph, note, size):
    run = paragraph.add_run(f" (Ghi chú: {note})")
    set_run_font(run, size)
    run.italic = True


def find_paragraph(paragraphs, needle):
    for paragraph in paragraphs:
        if needle in paragraph.text:
            return paragraph
    raise AssertionError(f"Could not find paragraph containing: {needle}")


before_text, before_media = package_text_signature(SOURCE)
BACKUP.parent.mkdir(parents=True, exist_ok=True)
copy2(SOURCE, BACKUP)
doc = Document(SOURCE)

# Re-apply the agreed typography without altering paragraph structure or wording.
for style_name, size in {
    "Normal": 12,
    "Report Title": 16,
    "Title": 16,
    "Heading 1": 13,
    "Heading 2": 12,
    "Heading 3": 12,
}.items():
    if style_name in doc.styles:
        style = doc.styles[style_name]
        style.font.name = FONT
        style.font.size = Pt(size)
        rpr = style._element.get_or_add_rPr()
        rfonts = rpr.rFonts
        if rfonts is None:
            rfonts = etree.SubElement(rpr, qn("w:rFonts"))
        for key in ("ascii", "hAnsi", "eastAsia", "cs"):
            rfonts.set(qn(f"w:{key}"), FONT)

for index, paragraph in enumerate(doc.paragraphs):
    if paragraph.style.name in ("Report Title", "Title"):
        size = 16
    elif paragraph.style.name == "Heading 1":
        size = 13
    elif index == 1 or index >= len(doc.paragraphs) - 4:
        size = 12
    else:
        size = 12
    for run in paragraph.runs:
        set_run_font(run, size)

for section in doc.sections:
    for part in (section.header, section.footer):
        for paragraph in part.paragraphs:
            for run in paragraph.runs:
                set_run_font(run, 9)

# Body notes: append only, preserving the user's exact original wording.
body_notes = {
    "Đối chiếu dữ liệu:": "cần thống nhất tiêu chí phân loại: tiếp tục dùng trên 1Office, tích hợp với SDG ERP hay phát triển mới trên ERP.",
    "Công tác bàn giao:": "cần xác nhận lại cách ghi “7 license Office 365”, tài khoản quản trị 1Office và tên chính thức của tài khoản ChatGPT thuộc HCNS.",
    "Tài khoản cá nhân nội bộ": "cần nêu rõ tài khoản nào được thay thế, thời điểm thực hiện và người chịu trách nhiệm bàn giao; đồng thời kiểm tra chính tả cụm “bộ phân IT”.",
    "Xây dựng hệ thống quản lý nguồn tài nguyên": "phạm vi SDG ERP cần được quản lý thống nhất; nên xác định rõ dữ liệu và nghiệp vụ ưu tiên để tránh hiểu là thay thế toàn bộ 1Office.",
    "Trợ lý tài liệu AI:": "cần xác định dữ liệu được phép sử dụng, phương thức kết nối, chi phí và bước kiểm tra/phê duyệt đầu ra.",
    "Mục tiêu phấn đấu vượt chỉ tiêu:": "chỉ nên dùng khái niệm “vượt chỉ tiêu” sau khi KPI và tiêu chí đánh giá được quản lý trực tiếp thống nhất.",
}
for needle, note in body_notes.items():
    append_note(find_paragraph(doc.paragraphs, needle), note, 12)

# Masthead and schedule table typography.
masthead = doc.tables[0]
for cell in masthead.rows[0].cells:
    for paragraph in cell.paragraphs:
        if "CỘNG HÒA" in paragraph.text:
            size = 10.5
        elif "Độc lập" in paragraph.text:
            size = 11.5
        else:
            size = 12
        for run in paragraph.runs:
            set_run_font(run, size)

schedule = doc.tables[1]
for row_index, row in enumerate(schedule.rows):
    size = 11 if row_index == 0 else 10.5
    for cell in row.cells:
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                set_run_font(run, size)

# Table notes are placed in the same cell as the point requiring clarification.
append_note(
    find_paragraph(schedule.cell(1, 2).paragraphs, "Đã nhận bàn giao các tài khoản"),
    "cần đối chiếu tên chính thức và số lượng tài khoản/license với biên bản bàn giao.",
    10.5,
)
append_note(
    find_paragraph(schedule.cell(2, 1).paragraphs, "tài sản vật lý-phi vật lý"),
    "cần xác định “tài sản phi vật lý” gồm license, tài khoản hay nhóm dữ liệu nào.",
    10.5,
)
append_note(
    find_paragraph(schedule.cell(2, 2).paragraphs, "01 báo cáo trước hết thử việc 5 ngày"),
    "cụm “trước hết thử việc 5 ngày” chưa rõ; cần xác định có phải là “trước khi hết thử việc 5 ngày” hay một mốc khác.",
    10.5,
)
append_note(
    schedule.cell(3, 0).paragraphs[0],
    "cần định nghĩa T là ngày kết thúc thử việc hay một mốc báo cáo khác.",
    10.5,
)
append_note(
    find_paragraph(schedule.cell(2, 2).paragraphs, "Hệ thống thông tin nhân sự-phòng ban"),
    "ô đầu ra này đang dùng chung cho ba mốc; cần xác nhận đây là kết quả chung hay cần tách theo từng giai đoạn.",
    10.5,
)
append_note(
    schedule.cell(4, 0).paragraphs[0],
    "cần thống nhất cách ghi mốc thời gian với các dòng phía trên.",
    10.5,
)
doc.save(OUTPUT)

after_text, after_media = package_text_signature(OUTPUT)
for part, original_nodes in before_text.items():
    assert part in after_text
    assert is_subsequence(original_nodes, after_text[part]), f"Original text changed in {part}"
assert before_media == after_media, "Embedded media changed"

check = Document(OUTPUT)
document_note_count = sum(node.count("(Ghi chú:") for node in after_text["word/document.xml"])
assert document_note_count == 12
assert len(check.tables) == 2 and len(check.tables[1].rows) == 5
print(f"Normalized and annotated report in place: {OUTPUT}")
print(f"Internal backup: {BACKUP}")
