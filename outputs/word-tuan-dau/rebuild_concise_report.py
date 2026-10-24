from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor, Twips


OUT = Path("docs/bao-cao/TONG-HOP-BAO-CAO-TUAN-DAU-VA-CONG-VIEC-CAN-LAM.docx")
BLACK = RGBColor(0, 0, 0)
FONT = "Times New Roman"


def set_font(run, size=12.5, bold=False, italic=False):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.font.color.rgb = BLACK
    run.bold = bold
    run.italic = italic


def set_cell_margins(cell, top=0, start=80, bottom=0, end=80):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_no_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "nil")
        borders.append(node)
    tbl_pr.append(borders)


def set_table_geometry(table, widths, indent=0):
    total = sum(widths)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = OxmlElement("w:tblInd")
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_pr.append(tbl_ind)
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Twips(width)
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")


def add_text(doc, text, *, bold_label=None, after=5, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.15
    if bold_label and text.startswith(bold_label):
        set_font(p.add_run(bold_label), bold=True)
        set_font(p.add_run(text[len(bold_label):]))
    else:
        set_font(p.add_run(text))
    return p


def add_item(doc, label, body):
    return add_text(doc, f"{label}: {body}", bold_label=f"{label}:", after=4)


def add_heading(doc, text):
    p = doc.add_paragraph(text, style="Heading 1")
    p.paragraph_format.keep_with_next = True
    return p


doc = Document()
section = doc.sections[0]
section.page_width = Cm(21)
section.page_height = Cm(29.7)
section.top_margin = Cm(2.0)
section.bottom_margin = Cm(2.0)
section.left_margin = Cm(2.5)
section.right_margin = Cm(2.0)
section.header_distance = Cm(1.0)
section.footer_distance = Cm(1.0)

# Resolved standard_business_brief with formal monochrome overrides:
# A4, Times New Roman, black-only hierarchy, compact body spacing.
normal = doc.styles["Normal"]
normal.font.name = FONT
normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
normal.font.size = Pt(12.5)
normal.font.color.rgb = BLACK
normal.paragraph_format.space_before = Pt(0)
normal.paragraph_format.space_after = Pt(5)
normal.paragraph_format.line_spacing = 1.15

for name, size, before, after in (
    ("Title", 16, 6, 8),
    ("Heading 1", 13, 10, 5),
    ("Heading 2", 12.5, 8, 4),
    ("Heading 3", 12.5, 6, 3),
):
    style = doc.styles[name]
    style.font.name = FONT
    style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = BLACK
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.line_spacing = 1.0

# Dedicated title style avoids the blue rule carried by Word's built-in Title
# style on this Office installation.
report_title = doc.styles.add_style("Report Title", WD_STYLE_TYPE.PARAGRAPH)
report_title.base_style = doc.styles["Normal"]
report_title.font.name = FONT
report_title._element.rPr.rFonts.set(qn("w:ascii"), FONT)
report_title._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
report_title._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
report_title.font.size = Pt(16)
report_title.font.bold = True
report_title.font.color.rgb = BLACK
report_title.paragraph_format.space_before = Pt(6)
report_title.paragraph_format.space_after = Pt(8)
report_title.paragraph_format.line_spacing = 1.0

# Quiet running header/footer in black only.
header_p = section.header.paragraphs[0]
header_p.text = "BÁO CÁO TUẦN ĐẦU THỬ VIỆC"
header_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
for run in header_p.runs:
    set_font(run, 9)

footer_p = section.footer.paragraphs[0]
footer_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
set_font(footer_p.add_run("Trang "), 9)
field = OxmlElement("w:fldSimple")
field.set(qn("w:instr"), "PAGE")
footer_p._p.append(field)

# Formal two-column masthead requested by the user.
masthead = doc.add_table(rows=1, cols=2)
set_table_geometry(masthead, [5600, 3600], indent=0)
set_no_borders(masthead)
left, right = masthead.rows[0].cells
for cell in (left, right):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    set_cell_margins(cell)

for cell, lines in (
    (left, (("CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM", 10.5, True), ("Độc lập - Tự do - Hạnh phúc", 11.5, True))),
    (right, (("TẬP ĐOÀN SAO ĐỎ", 12, True), ("Số: …/BC", 12, False))),
):
    for index, (text, size, bold) in enumerate(lines):
        p = cell.paragraphs[0] if index == 0 else cell.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(3 if index == 0 else 8)
        p.paragraph_format.keep_with_next = True
        set_font(p.add_run(text), size, bold=bold)

title = doc.add_paragraph(style="Report Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title_ppr = title._p.get_or_add_pPr()
title_border = OxmlElement("w:pBdr")
for edge in ("top", "left", "bottom", "right", "between"):
    node = OxmlElement(f"w:{edge}")
    node.set(qn("w:val"), "nil")
    title_border.append(node)
title_ppr.append(title_border)
set_font(title.add_run("BÁO CÁO TUẦN ĐẦU THỬ VIỆC\nVÀ KẾ HOẠCH CÔNG VIỆC TIẾP THEO"), 16, bold=True)

meta = doc.add_paragraph()
meta.paragraph_format.space_after = Pt(7)
meta.paragraph_format.line_spacing = 1.1
for line in (
    "Người báo cáo: [Họ và tên]",
    "Vị trí: Nhân viên IT - Phòng Hành chính - Nhân sự",
    "Kính gửi: [Họ tên/chức danh người nhận]",
    "Kỳ báo cáo: Từ ngày […] đến ngày […]",
):
    set_font(meta.add_run(line + "\n"), 12)

add_heading(doc, "1. Mục đích báo cáo")
add_text(
    doc,
    "Báo cáo tóm tắt công việc đã thực hiện trong tuần đầu thử việc, kết quả bước đầu và kế hoạch tiếp theo để thống nhất với quản lý trực tiếp.",
)

add_heading(doc, "2. Công việc đã thực hiện và kết quả")
add_item(doc, "Khảo sát hiện trạng", "Đọc tài liệu nội bộ, cẩm nang hội nhập và tìm hiểu cách công ty đang sử dụng 1Office.")
add_item(doc, "Đối chiếu dữ liệu", "Xác định 1Office là nguồn chuẩn về nhân sự và cơ cấu tổ chức. Snapshot ngày 21/10/2026 cho thấy 18/18 đơn vị trên ERP khớp tên, ID nguồn và quan hệ cha; chưa xem đây là kết quả nghiệm thu toàn bộ tính năng.")
add_item(doc, "Xây dựng bản mẫu", "Đã có bản mẫu cơ cấu phòng ban, dữ liệu tài sản và mô hình văn phòng 3D để phục vụ khảo sát.")
add_item(doc, "Nghiên cứu tài liệu AI", "Xây dựng hướng thử nghiệm nhận ảnh hóa đơn, trích xuất thông tin, cho người dùng kiểm tra và tạo file Word theo mẫu. Đây là bản thử nghiệm, chưa phải quy trình vận hành chính thức.")

add_heading(doc, "3. Hướng đề xuất")
add_item(doc, "Quản lý tài sản gắn với không gian 3D", "Tra cứu thiết bị, người sử dụng, vị trí, tình trạng và lịch sử xử lý trên một khu vực thử nghiệm.")
add_item(doc, "Chuẩn hóa dữ liệu nền", "Dùng dữ liệu nhân sự và cơ cấu từ 1Office; ERP lưu bản sao phục vụ các nghiệp vụ riêng và cần có cơ chế đối soát.")
add_item(doc, "Trợ lý tài liệu AI", "Hỗ trợ đọc ảnh/PDF, lấy thông tin, soạn nội dung và xuất Word theo mẫu; người dùng phải kiểm tra trước khi sử dụng.")

add_heading(doc, "4. Mẫu và tư liệu minh họa khi trình bày")
add_text(doc, "Khi báo cáo, dự kiến trình bày ngắn các nội dung sau:")
add_item(doc, "Mẫu trực tiếp", "Giao diện cơ cấu phòng ban, quản lý tài sản, mô hình 3D và luồng tạo hồ sơ thanh toán từ ảnh sang file Word.")
add_item(doc, "Hình ảnh hoặc video", "Dùng ảnh chụp màn hình hoặc video thao tác ngắn để minh họa chức năng và kết quả đầu ra khi không thuận tiện trình diễn trực tiếp.")

add_heading(doc, "5. Công việc tiếp theo")
add_item(doc, "Thống nhất mục tiêu", "Chốt KPI, thời hạn, đầu ra, người đánh giá và đầu mối phối hợp.")
add_item(doc, "Hoàn thiện cơ cấu phòng ban", "Kiểm thử quyền truy cập, đối chiếu dữ liệu và quy trình nhập snapshot.")
add_item(doc, "Thử nghiệm tài sản và 3D", "Chọn một nhóm thiết bị hoặc một khu vực nhỏ, xác nhận dữ liệu và lấy phản hồi người dùng.")
add_item(doc, "Thử nghiệm tài liệu AI", "Dùng một bộ hồ sơ thanh toán phù hợp để kiểm tra quá trình nhập ảnh, rà soát dữ liệu và xuất Word.")
add_item(doc, "Nghiên cứu tích hợp 1Office", "Xác minh khả năng cung cấp API/webhook, cách cập nhật dữ liệu và xử lý khi mất kết nối.")

add_heading(doc, "6. Nội dung cần thống nhất")
add_text(
    doc,
    "Đề nghị quản lý trực tiếp thống nhất KPI, thứ tự ưu tiên, phạm vi thử nghiệm và đầu mối phối hợp. Các hướng nêu trên mới là đề xuất phục vụ trao đổi, chưa phải kế hoạch triển khai đã được phê duyệt.",
)

add_text(doc, "Tôi kính báo cáo và đề nghị quản lý trực tiếp góp ý để hoàn thiện kế hoạch công việc.", after=10)

for text, bold, italic, before, after in (
    ("………, ngày …… tháng …… năm ……", False, True, 4, 4),
    ("NGƯỜI BÁO CÁO", True, False, 0, 2),
    ("(Ký và ghi rõ họ tên)", False, True, 0, 28),
    ("[Họ và tên]", True, False, 0, 0),
):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.keep_with_next = text != "[Họ và tên]"
    set_font(p.add_run(text), 12, bold=bold, italic=italic)

doc.core_properties.title = "Báo cáo tuần đầu thử việc và kế hoạch công việc tiếp theo"
doc.core_properties.subject = "Bản rút gọn, trình bày đơn sắc"
doc.core_properties.author = ""
doc.save(OUT)

# Structural checks before Microsoft Word render.
check = Document(OUT)
content = "\n".join(p.text for p in check.paragraphs) + "\n" + "\n".join(c.text for t in check.tables for r in t.rows for c in r.cells)
assert len(check.tables) == 1
assert "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM" in content
assert "Mẫu và tư liệu minh họa khi trình bày" in content
assert "Hình ảnh hoặc video" in content
assert sum(p.style.name == "Heading 1" for p in check.paragraphs) == 6
assert all((run.font.color.rgb in (None, BLACK)) for p in check.paragraphs for run in p.runs)
print(f"Rebuilt and checked: {OUT}")
