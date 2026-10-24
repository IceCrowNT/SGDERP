from pathlib import Path
from docx import Document
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
p=Path('docs/bao-cao/DANH-SACH-TAI-SAN-IT-VA-GHI-CHU.docx')
d=Document(p)
# Preserve the summary table, replace everything after it with a plain per-device list.
summary=d.tables[0]
items=[(r.cells[1].text,int(r.cells[2].text)) for r in summary.rows[1:-1]]
node=summary._tbl.getnext()
while node is not None and node.tag != qn('w:sectPr'):
 nxt=node.getnext(); node.getparent().remove(node); node=nxt
d.add_page_break()
d.add_paragraph('DANH SÁCH TỪNG THIẾT BỊ', 'Title')
d.add_paragraph('Danh sách tạm theo số lượng trên ảnh bàn giao; chưa xác nhận từng thiết bị qua kiểm kê. STT chỉ dùng để ghi chú, không phải mã tài sản.')
t=d.add_table(rows=1,cols=3); t.autofit=False
widths=[600,4500,4260]
from docx.shared import Twips
for c,w in zip(t.columns,widths): c.width=Twips(w)
for c,txt in zip(t.rows[0].cells,['STT','Thiết bị','Ghi chú']): c.text=txt
n=0
for name,count in items:
 for _ in range(count):
  n+=1
  for c,txt in zip(t.add_row().cells,[str(n),name,'']): c.text=txt
for i,row in enumerate(t.rows):
 pr=row._tr.get_or_add_trPr(); pr.append(OxmlElement('w:cantSplit'))
 if i==0: pr.append(OxmlElement('w:tblHeader'))
 for c,w in zip(row.cells,widths):
  c.width=Twips(w)
  cp=c._tc.get_or_add_tcPr(); borders=OxmlElement('w:tcBorders')
  for edge in ['top','left','bottom','right']:
   e=OxmlElement('w:'+edge); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'AAAAAA'); borders.append(e)
  cp.append(borders)
  for para in c.paragraphs:
   para.paragraph_format.space_before=Pt(4); para.paragraph_format.space_after=Pt(4); para.paragraph_format.line_spacing=1
   for r in para.runs: r.font.size=Pt(10); r.bold=i==0
assert n==90
d.save(p)
check=Document(p); assert len(check.tables[-1].rows)==91
print('Updated: summary + 90 device rows, blank notes.')
