from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt
from lxml import etree


SOURCE = Path("docs/bao-cao/TONG-HOP-BAO-CAO-TUAN-DAU-VA-CONG-VIEC-CAN-LAM.docx")
OUTPUT = Path("docs/bao-cao/TONG-HOP-BAO-CAO-TUAN-DAU-VA-CONG-VIEC-CAN-LAM-DONG-BO-FONT.docx")
FONT = "Times New Roman"
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def package_text_signature(path: Path):
    signature = {}
    with ZipFile(path) as package:
        names = package.namelist()
        selected = [
            name
            for name in names
            if name == "word/document.xml"
            or name.startswith("word/header")
            or name.startswith("word/footer")
            or name in {
                "word/comments.xml",
                "word/footnotes.xml",
                "word/endnotes.xml",
            }
        ]
        for name in sorted(selected):
            root = etree.fromstring(package.read(name))
            signature[name] = tuple(root.xpath("//w:t/text()", namespaces={"w": W_NS}))
        media = {
            name: package.getinfo(name).file_size
            for name in names
            if name.startswith("word/media/")
        }
    return signature, media


def set_run_font(run, size):
    run.font.name = FONT
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = etree.SubElement(rpr, qn("w:rFonts"))
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{key}"), FONT)
    run.font.size = Pt(size)


def set_style_font(style, size):
    style.font.name = FONT
    rpr = style._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = etree.SubElement(rpr, qn("w:rFonts"))
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{key}"), FONT)
    style.font.size = Pt(size)


before_text, before_media = package_text_signature(SOURCE)
doc = Document(SOURCE)

# Style-level defaults for new typing in the normalized copy.
style_sizes = {
    "Normal": 12,
    "Report Title": 16,
    "Title": 16,
    "Heading 1": 13,
    "Heading 2": 12,
    "Heading 3": 12,
}
for style_name, size in style_sizes.items():
    if style_name in doc.styles:
        set_style_font(doc.styles[style_name], size)

# Main document paragraphs. Bold/italic, alignment, spacing and wording remain unchanged.
for index, paragraph in enumerate(doc.paragraphs):
    style_name = paragraph.style.name
    if style_name in ("Report Title", "Title"):
        size = 16
    elif style_name == "Heading 1":
        size = 13
    elif style_name in ("Heading 2", "Heading 3"):
        size = 12
    elif index == 1 or index >= len(doc.paragraphs) - 4:
        size = 12
    else:
        size = 12
    for run in paragraph.runs:
        set_run_font(run, size)

# Two-column formal masthead.
if doc.tables:
    masthead = doc.tables[0]
    for cell in masthead.rows[0].cells:
        for paragraph_index, paragraph in enumerate(cell.paragraphs):
            if "CỘNG HÒA" in paragraph.text:
                size = 10.5
            elif "Độc lập" in paragraph.text:
                size = 11.5
            else:
                size = 12
            for run in paragraph.runs:
                set_run_font(run, size)

# Data tables: one consistent body size and a slightly larger header row.
for table in doc.tables[1:]:
    for row_index, row in enumerate(table.rows):
        size = 11 if row_index == 0 else 10.5
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    set_run_font(run, size)

# Running headers and footers.
for section in doc.sections:
    for part in (section.header, section.footer):
        for paragraph in part.paragraphs:
            for run in paragraph.runs:
                set_run_font(run, 9)
        for table in part.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        for run in paragraph.runs:
                            set_run_font(run, 9)

doc.save(OUTPUT)

after_text, after_media = package_text_signature(OUTPUT)
assert before_text == after_text, "Text content changed during font normalization"
assert before_media == after_media, "Embedded media changed during font normalization"

check = Document(OUTPUT)
assert len(check.paragraphs) == len(doc.paragraphs)
assert len(check.tables) == len(doc.tables)
assert len(check.tables) == 2
assert len(check.tables[1].rows) == 5
print(f"Created normalized copy without text changes: {OUTPUT}")
