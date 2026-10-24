from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

path = Path('docs/bao-cao/TONG-HOP-BAO-CAO-TUAN-DAU-VA-CONG-VIEC-CAN-LAM.docx')
doc = Document(path)
# Preserve the existing report and A4 layout. Named override:
# formal national masthead, centered black title, and signature block.
title = doc.paragraphs[0]
for text, size, after in [
    ('CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM', 12, 3),
    ('Độc lập - Tự do - Hạnh phúc', 12, 10),
]:
    p = title.insert_paragraph_before(text)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.keep_with_next = True
    for r in p.runs:
        r.bold = True
        r.font.size = Pt(size)
        r.font.color.rgb = RGBColor(0, 0, 0)

title.alignment = WD_ALIGN_PARAGRAPH.CENTER
for r in title.runs:
    r.font.size = Pt(17)
    r.font.color.rgb = RGBColor(0, 0, 0)

for p in doc.paragraphs:
    if p.text.startswith('Bản nháp báo cáo công việc'):
        p.text = 'Bản dự thảo trình quản lý trực tiếp • Cập nhật ngày 22/10/2026'
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.size = Pt(10)
            r.italic = True
    elif p.text.startswith('Người báo cáo:'):
        p.text = 'Người báo cáo: [Họ và tên]\nVị trí: Nhân viên IT – Phòng Hành chính – Nhân sự\nKính gửi: [Họ tên/chức danh người nhận]\nKỳ báo cáo: Từ ngày […] đến ngày […]'
    elif p.text.startswith('Mục tiêu gần nhất là'):
        p.text = 'Tôi báo cáo kết quả tìm hiểu, khảo sát và xây dựng bản mẫu trong tuần đầu thử việc tại Sao Đỏ, đồng thời đề xuất công việc tiếp theo để thống nhất với quản lý trực tiếp.'
    elif p.text.startswith('Người phụ trách là nhân viên IT'):
        p.text = 'Với vị trí nhân viên IT thuộc phòng Hành chính – Nhân sự (HCNS), tôi tập trung tìm hiểu hiện trạng và nhu cầu số hóa. Do thông tin bàn giao ban đầu còn hạn chế, các nội dung được tổng hợp từ tài liệu nội bộ, khảo sát hệ thống và trao đổi với người phụ trách.'

outcomes = {
    'Ưu tiên 1': ' Đầu ra: danh sách kết quả đã có và các hạng mục còn cần hoàn thiện.',
    'Ưu tiên 2': ' Đầu ra: bảng mục tiêu có chỉ số đo, người đánh giá và thời hạn được thống nhất.',
    'Ưu tiên 3': ' Đầu ra: biên bản kiểm thử quyền theo vai trò, đối chiếu dữ liệu và hướng dẫn nhập snapshot.',
    'Nghiên cứu sâu 1Office:': ' Đầu ra: danh sách nhu cầu, tính năng có thể tận dụng và vấn đề cần xác minh với đầu mối nghiệp vụ.',
    'Thử quản lý từng tài sản:': ' Đầu ra: danh sách thiết bị thử nghiệm được đầu mối quản lý xác nhận và phạm vi quyền tương ứng.',
    'Đánh giá và thử mô hình 3D:': ' Đầu ra: một khu vực mẫu gắn đúng vị trí thiết bị và ghi nhận phản hồi người dùng.',
    'Thử một nghiệp vụ AI:': ' Đầu ra: mẫu tài liệu đã được người dùng kiểm tra và ghi nhận thời gian chỉnh sửa; chỉ dùng dữ liệu được phép, đã ẩn danh đầy đủ.',
    'Thiết kế đồng bộ 1Office–ERP:': ' Đầu ra: phương án đồng bộ một chiều, danh sách API cần xác minh và cách đối soát khi gián đoạn.',
}
for p in doc.paragraphs:
    for prefix, tail in outcomes.items():
        if p.text.startswith(prefix):
            p.add_run(tail)

heading6 = next(p for p in doc.paragraphs if p.text.startswith('6. Những nội dung'))
p = heading6.insert_paragraph_before('Thứ tự thực hiện đề xuất: tuần tiếp theo ưu tiên rà soát kết quả, thống nhất KPI, hoàn thiện cơ cấu và khảo sát nhu cầu. Thử nghiệm tài sản, 3D, AI và xây dựng tích hợp thực hiện sau khi được giao phạm vi cụ thể; thời hạn chi tiết thống nhất với quản lý trực tiếp.')
p.paragraph_format.space_after = Pt(8)

for p in doc.paragraphs:
    if p.text.startswith('7. Tài liệu tham chiếu'):
        p.paragraph_format.page_break_before = False
    p.paragraph_format.widow_control = True

p = doc.add_paragraph('Tôi kính báo cáo và đề nghị quản lý trực tiếp góp ý, thống nhất thứ tự ưu tiên, đầu mối phối hợp và tiêu chí đánh giá cho giai đoạn tiếp theo.')
p.paragraph_format.keep_with_next = True
for text, bold, italic, before, after in [
    ('………, ngày …… tháng …… năm ……', False, True, 12, 6),
    ('NGƯỜI BÁO CÁO', True, False, 0, 3),
    ('(Ký và ghi rõ họ tên)', False, True, 0, 36),
    ('[Họ và tên]', True, False, 0, 0),
]:
    p = doc.add_paragraph(text)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.keep_with_next = text != '[Họ và tên]'
    for r in p.runs:
        r.bold, r.italic = bold, italic

doc.core_properties.subject = 'Báo cáo kết quả tuần đầu thử việc và kế hoạch công việc đề xuất'
doc.save(path)
check = Document(path)
content = '\n'.join(p.text for p in check.paragraphs)
assert 'CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM' in content
assert 'Độc lập - Tự do - Hạnh phúc' in content
assert sum(p.style.name == 'Heading 1' for p in check.paragraphs) == 7
assert content.count('Đầu ra:') == 8
assert 'NGƯỜI BÁO CÁO' in content
print(f'Updated and structurally checked: {path}')
