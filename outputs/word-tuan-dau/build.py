from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pathlib import Path
# standard_business_brief; named overrides: A4, Arial, SDG burgundy headings.
d=Document(); s=d.sections[0]; s.page_width=Inches(8.2677);s.page_height=Inches(11.6929)
for side in ['top_margin','bottom_margin','left_margin','right_margin']:setattr(s,side,Inches(1))
s.header_distance=s.footer_distance=Inches(.492)
for name,size,before,after,color in [('Normal',11,0,6,'24262B'),('Title',22,0,12,'A32632'),('Heading 1',16,16,8,'A32632'),('Heading 2',13,12,6,'A32632'),('Heading 3',12,8,4,'24262B')]:
 st=d.styles[name];st.font.name='Arial';st.font.size=Pt(size);st.font.color.rgb=RGBColor.from_string(color)
 st.paragraph_format.space_before=Pt(before);st.paragraph_format.space_after=Pt(after);st.paragraph_format.line_spacing=1.10
 if name!='Normal':st.font.bold=True;st.paragraph_format.keep_with_next=True
h=s.header.paragraphs[0];h.text='SAO ĐỎ GROUP  |  IT – HÀNH CHÍNH NHÂN SỰ';h.style='Normal'
for r in h.runs:r.font.size=Pt(9);r.font.color.rgb=RGBColor.from_string('666666')
f=s.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=f.add_run('Báo cáo công việc tuần đầu  •  Trang ');r.font.size=Pt(9)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld)
def p(t):d.add_paragraph(t)
def h1(t):d.add_heading(t,1)
def item(title,body):
 a=d.add_paragraph();a.add_run(title+' ').bold=True;a.add_run(body)
def page():d.add_page_break()
d.add_paragraph('BÁO CÁO TUẦN ĐẦU THỬ VIỆC\nVÀ CÁC CÔNG VIỆC CẦN LÀM',style='Title')
p('Bản nháp báo cáo công việc • Cập nhật ngày 22/10/2026')
p('Người báo cáo: [Cần bổ sung]\nNgười nhận: [Cần bổ sung]\nTuần báo cáo: [Từ ngày … đến ngày …]')
h1('1. Mục tiêu và bối cảnh')
p('Mục tiêu gần nhất là hoàn thiện báo cáo tuần đầu thử việc tại Sao Đỏ: trình bày những gì đã tìm hiểu, kết quả bước đầu và công việc dự kiến tiếp theo để trao đổi với quản lý trực tiếp.')
p('Người phụ trách là nhân viên IT mới thuộc phòng Hành chính – Nhân sự (HCNS), được giao nhiệm vụ số hóa. Thông tin bàn giao còn ít; phần lớn phải tự tìm hiểu tài liệu, khảo sát hệ thống và xác nhận với các phòng ban.')
p('1Office đang hoạt động ổn định và tiếp tục là hệ thống làm việc chính. SDG ERP được định hướng đồng hành, bổ sung nhu cầu đặc thù của Sao Đỏ. Các tính năng ERP hiện có vẫn là bản nháp phục vụ khảo sát.')
h1('2. Trọng tâm công việc tuần đầu')
item('Tìm hiểu tổ chức và cách phối hợp:', 'Nắm cơ cấu phòng ban, vai trò của IT trong HCNS và các đầu mối nghiệp vụ cần làm việc cùng.')
item('Khảo sát hệ thống đang sử dụng:', 'Tìm hiểu quy trình, dữ liệu và các tính năng 1Office để xác định khả năng tận dụng trước khi phát triển thêm.')
item('Đối chiếu và xây nền dữ liệu:', 'Kiểm tra nguồn nhân sự, cơ cấu tổ chức; lấy 1Office làm chuẩn áp dụng cho ERP.')
item('Xây dựng và đánh giá bản mẫu:', 'Thử nghiệm cơ cấu phòng ban, cơ sở dữ liệu tài sản và mô hình 3D để làm rõ nhu cầu thực tế.')
item('Đề xuất mục tiêu tiếp theo:', 'Chuẩn bị 3 hướng phát triển và các đầu ra để trao đổi, thống nhất với quản lý trực tiếp.')
page()
h1('3. Công việc đã thực hiện và kết quả bước đầu')
item('Tìm hiểu hiện trạng:', 'Đã đọc tài liệu nội bộ, cẩm nang hội nhập và khảo sát cách công ty sử dụng 1Office, các quy trình và dữ liệu đang có.')
item('Đối chiếu dữ liệu nền:', 'Lấy 1Office làm nguồn chuẩn về nhân sự và cơ cấu tổ chức. Theo snapshot ngày 21/10/2026, 18/18 đơn vị trên ERP khớp tên, ID nguồn và quan hệ cha; nguồn có thêm nút công ty. Kết quả này chưa đồng nghĩa nghiệm thu toàn bộ tính năng.')
item('Xây dựng nền tảng thử nghiệm:', 'Đã có bản mẫu cơ cấu phòng ban, cơ sở dữ liệu tài sản và mô hình văn phòng 3D. Chưa triển khai tích hợp API/webhook hoặc tính năng AI thực tế.')
item('Ghi nhận điểm cần xác minh:', 'Tiếp tục làm rõ trạng thái nhân sự, quyền truy cập, đầu mối phối hợp và quy trình nghiệp vụ. Không kết luận hệ thống có lỗi chỉ từ các dấu hiệu chưa được xác nhận.')
h1('4. Ba ý tưởng phát triển đã thảo luận')
d.add_heading('4.1. Quản lý trực quan bằng mô hình 3D',2)
p('Hiển thị nhân sự và thiết bị theo vị trí trong tòa nhà. Khi chọn một vị trí hoặc thiết bị, có thể xem người sử dụng, tình trạng, lịch sử sửa chữa và công việc liên quan. IoT là khả năng mở rộng về sau, chưa thuộc phạm vi triển khai đã chốt.')
p('Cần đánh giá mô hình hiện tại về độ chính xác với không gian thực, khả năng mở rộng, tính thẩm mỹ và mức độ dễ thao tác. Bắt đầu với một khu vực nhỏ để kiểm chứng giá trị sử dụng.')
d.add_heading('4.2. Chuẩn hóa và đồng bộ dữ liệu nền',2)
p('Thống nhất mã nhân viên, mã phòng ban và quan hệ tổ chức theo nguồn 1Office. ERP giữ bản sao cục bộ để phục vụ truy vấn nhanh và liên kết nghiệp vụ riêng, tránh duy trì hai bộ dữ liệu nhân sự độc lập.')
p('Hướng nghiên cứu gồm cập nhật định kỳ, webhook nếu được hỗ trợ, đối soát và xử lý gián đoạn. Không mặc định webhook bảo đảm dữ liệu luôn tức thời hoặc tự lưu lại mọi sự kiện khi một bên mất kết nối.')
d.add_heading('4.3. AI hỗ trợ nghiệp vụ nội bộ',2)
p('Xây sẵn bộ kỹ năng, quy tắc và biểu mẫu. Ví dụ: ảnh hóa đơn → trích thông tin → người dùng xác nhận, bổ sung mục đích → tạo bản nháp đề nghị thanh toán và tờ trình → rà soát, xuất Word.')
p('Mỗi phòng có tài khoản AI riêng theo thông tin đã trao đổi. Mong muốn là sử dụng riêng để tránh dồn chi phí; khả năng tích hợp các tài khoản hiện có vào web cần kiểm chứng. AI không tự tạo thông tin nghiệm thu, phê duyệt hoặc dữ kiện không có trong nguồn.')
page()
h1('5. Công việc cần làm tiếp')
item('Ưu tiên 1 — Chốt kết quả công việc:', 'Rà soát kết quả khảo sát, các dữ liệu đã đối chiếu và những phần còn dở; xác định đầu ra cần hoàn thiện trong giai đoạn tiếp theo.')
item('Ưu tiên 2 — Chốt mục tiêu với quản lý:', 'Thống nhất 3 KPI, thời hạn, đầu ra và tiêu chí đánh giá. Các mục tiêu trong bản trình bày hiện là đề xuất, chưa phải KPI đã được phê duyệt.')
item('Ưu tiên 3 — Hoàn thiện cơ cấu phòng ban:', 'Giữ nguồn chuẩn từ 1Office; giải quyết các phần kiểm thử quyền, đối chiếu và nhập snapshot còn dở trước khi nghiệm thu.')
item('Nghiên cứu sâu 1Office:', 'Xác định tính năng sẵn có nên tận dụng, nội dung có thể tối ưu và nhu cầu phù hợp phát triển thêm trên ERP. Chỉ khảo sát khi được cấp lại quyền truy cập trong phạm vi cụ thể.')
item('Thử quản lý từng tài sản:', 'Ảnh biên bản bàn giao chỉ là dữ liệu tham khảo. Cần bổ sung danh mục từng thiết bị, vị trí, người sử dụng và tình trạng thực tế. Chốt quyền theo phòng ban hoặc người được giao; không mặc định IT quản lý toàn bộ tài sản công ty.')
item('Đánh giá và thử mô hình 3D:', 'Kiểm tra vị trí thật, cách gán nhân sự/tài sản, khả năng mở rộng và trải nghiệm thao tác trước khi mở rộng phạm vi.')
item('Thử một nghiệp vụ AI:', 'Chọn bộ mẫu hồ sơ thanh toán, xác định trường thông tin bắt buộc, thử bằng dữ liệu phù hợp và để người dùng kiểm tra trước khi xuất tài liệu. Làm rõ cách tích hợp và chi phí trước khi triển khai.')
item('Thiết kế đồng bộ 1Office–ERP:', 'Nghiên cứu bản sao dữ liệu cục bộ, cập nhật định kỳ, webhook, đối soát và phục hồi sau gián đoạn. Đây là hướng nghiên cứu, chưa phải tính năng đang vận hành.')
h1('6. Những nội dung cần xác nhận trước khi chốt')
p('Cần thống nhất 3 KPI với quản lý, thời hạn và tiêu chí đánh giá; xác định đầu mối HCNS và các phòng phối hợp; chốt phạm vi thử nghiệm tài sản, 3D và AI được giao.')
p('Ưu tiên trước mắt là hoàn thiện báo cáo và thống nhất công việc tiếp theo. Không đặt yêu cầu hoàn thành cả ba ý tưởng ngay trong tuần đầu thử việc.')
h1('7. Tài liệu tham chiếu')
p('Cẩm nang hội nhập Tập đoàn Sao Đỏ, slide 20 (các nội dung về mục tiêu công việc, cơ cấu tổ chức và phối hợp); tài liệu bối cảnh và định hướng số hóa; tài liệu đồng bộ dữ liệu nền và phân quyền; checklist HR và chuyển đổi số; các trao đổi với người phụ trách.')
p('Tài liệu này tập trung vào kết quả công việc và kế hoạch tiếp theo. Các ý tưởng phát triển là đề xuất để trao đổi; không đồng nghĩa đã được phê duyệt triển khai.')
d.core_properties.title='Báo cáo tuần đầu thử việc và các công việc cần làm';d.core_properties.author='';d.core_properties.subject='Tổng hợp nội dung đã thảo luận; chờ xác nhận'
out=Path('docs/bao-cao/TONG-HOP-BAO-CAO-TUAN-DAU-VA-CONG-VIEC-CAN-LAM.docx');d.save(out);print(out)
