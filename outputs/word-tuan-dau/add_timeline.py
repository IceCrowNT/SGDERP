from docx import Document
from docx.shared import Pt, Cm
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path
path=Path('docs/bao-cao/TONG-HOP-BAO-CAO-TUAN-DAU-VA-CONG-VIEC-CAN-LAM.docx')
doc=Document(path)
start=next(p for p in doc.paragraphs if p.text=='5. Công việc tiếp theo')
end=next(p for p in doc.paragraphs if p.text=='6. Nội dung cần thống nhất')
start.runs[0].text='5. Tiến độ dự kiến và kết quả đối chiếu'
for r in start.runs[1:]: r.text=''
node=start._p.getnext()
while node is not None and node is not end._p:
 nxt=node.getnext(); node.getparent().remove(node); node=nxt

def para(text):
 p=doc.add_paragraph(text)
 p.paragraph_format.space_after=Pt(5)
 end._p.addprevious(p._p)
 return p
para('Các mốc tính từ ngày bắt đầu thử việc. T là ngày kết thúc thử việc; báo cáo tổng kết hoàn thành chậm nhất T−5 ngày. Ngày cụ thể và chỉ tiêu được chốt với quản lý trực tiếp; nếu thời gian thử việc ngắn hơn, điều chỉnh phạm vi để bảo đảm hạn báo cáo.')
rows=[
('Mốc dự kiến','Công việc trọng tâm','Đầu ra để đối chiếu'),
('Tuần 1','Khảo sát hiện trạng, rà soát 1Office và xác định hướng phát triển.','01 báo cáo tuần đầu; 03 hướng đề xuất: dữ liệu nền, tài sản/3D và tài liệu AI.'),
('Tuần 2–4 / cuối tháng 1','Ưu tiên cơ cấu phòng ban và cập nhật dữ liệu bằng file; thống nhất phạm vi thử nghiệm tài sản.','01 luồng nhập cơ cấu có xem trước thay đổi và xác nhận; 01 lần nhập thử có kết quả đối chiếu; danh sách trường tài sản cần quản lý.'),
('Tháng 2 / trước T−10 ngày, nếu còn trong kỳ thử việc','Thử nghiệm tài sản và vị trí tại 01 khu vực nhỏ; hoàn thiện theo phản hồi. Khảo sát API/webhook.','Danh sách từng tài sản trong phạm vi đã thống nhất; minh họa tra cứu trên 3D từ vị trí xác nhận; 01 bản tổng hợp phản hồi và đánh giá phương án tích hợp.'),
('T−10 đến T−5 ngày','Tổng hợp minh chứng, so sánh chỉ tiêu và kết quả; hoàn thiện báo cáo cuối kỳ.','01 báo cáo trước hết thử việc 5 ngày, kèm ảnh/video hoặc bản chạy thử; phân loại đạt, vượt, chưa đạt và kế hoạch tiếp theo.'),
]
t=doc.add_table(rows=0, cols=3)
t.autofit=False
widths=[Cm(3),Cm(6),Cm(7)]
for col,w in zip(t.columns,widths): col.width=w
for ri,row in enumerate(rows):
 cells=t.add_row().cells
 trpr=t.rows[-1]._tr.get_or_add_trPr()
 trpr.append(OxmlElement('w:cantSplit'))
 if ri==0: trpr.append(OxmlElement('w:tblHeader'))
 for cell,w,text in zip(cells,widths,row):
  cell.width=w; cell.text=text
  for p in cell.paragraphs:
   p.paragraph_format.space_after=Pt(5); p.paragraph_format.space_before=Pt(5); p.paragraph_format.line_spacing=1.05
   for r in p.runs:
    r.font.name='Times New Roman'; r.font.size=Pt(11); r.bold=(ri==0)
  borders=OxmlElement('w:tcBorders')
  for edge in ['top','left','bottom','right']:
   el=OxmlElement('w:'+edge); el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'4'); el.set(qn('w:color'),'999999'); borders.append(el)
  cell._tc.get_or_add_tcPr().append(borders)
end._p.addprevious(t._tbl)
para('Mục tiêu phấn đấu vượt chỉ tiêu: khi hoàn thành phần ưu tiên và còn thời gian, bổ sung 01 bộ hồ sơ thanh toán xuất Word theo mẫu và thử nghiệm AI nếu có kết nối, ngân sách; hoặc mở rộng thử nghiệm tài sản sang thêm 01 khu vực. Chọn theo nhu cầu thực tế, không làm chậm đầu ra đã thống nhất.')
para('Cách đánh giá: mỗi đầu ra có minh chứng và phản hồi của đầu mối sử dụng. Ghi nhận vượt chỉ tiêu khi có đầu ra bổ sung đã kiểm chứng hoặc hoàn thành sớm mà vẫn đáp ứng phạm vi và chất lượng; không tính riêng số màn hình hay ý tưởng. Phần đã làm ở tuần 1 được ghi nhận là kết quả nền, không tính lại thành thành tích mới.')
doc.save(path)
check=Document(path)
assert any('T−5' in p.text for p in check.paragraphs)
assert len(check.tables[-1].rows)==5
print('Updated timeline; DOCX reopened and table structure checked.')
