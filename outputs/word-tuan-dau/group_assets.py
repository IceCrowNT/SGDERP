from pathlib import Path
from docx import Document
from docx.shared import Pt, Twips
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
p=Path('docs/bao-cao/DANH-SACH-TAI-SAN-IT-VA-GHI-CHU.docx')
d=Document(p)
summary=d.tables[0]
items=[(r.cells[1].text,int(r.cells[2].text)) for r in summary.rows[1:-1]]
# Replace both old tables and intervening list title with one grouped table.
node=summary._tbl
while node is not None and node.tag != qn('w:sectPr'):
 nxt=node.getnext(); node.getparent().remove(node); node=nxt
for paragraph in d.paragraphs:
 if paragraph.text.startswith('Người giao ghi trên ảnh:'):
  paragraph.add_run(' Danh sách gồm 11 nhóm, tổng 90 thiết bị. STT chỉ dùng để ghi chú, chưa phải mã tài sản đã kiểm kê.')
t=d.add_table(rows=0,cols=3); t.autofit=False
widths=[600,4500,4260]
for col,w in zip(t.columns,widths): col.width=Twips(w)
pr=t._tbl.tblPr
pr.find(qn('w:tblW')).set(qn('w:w'),'9360')
pr.find(qn('w:tblW')).set(qn('w:type'),'dxa')
ind=OxmlElement('w:tblInd'); ind.set(qn('w:w'),'120'); ind.set(qn('w:type'),'dxa'); pr.append(ind)
def row(values,kind='detail'):
 r=t.add_row(); rp=r._tr.get_or_add_trPr(); rp.append(OxmlElement('w:cantSplit'))
 if kind=='header': rp.append(OxmlElement('w:tblHeader'))
 for c,w,txt in zip(r.cells,widths,values):
  c.width=Twips(w); c.text=txt
  cp=c._tc.get_or_add_tcPr()
  borders=OxmlElement('w:tcBorders')
  for edge in ['top','left','bottom','right']:
   e=OxmlElement('w:'+edge); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'AAAAAA'); borders.append(e)
  cp.append(borders)
  margins=OxmlElement('w:tcMar')
  for edge,val in [('top',80),('bottom',80),('start',120),('end',120)]:
   e=OxmlElement('w:'+edge); e.set(qn('w:w'),str(val)); e.set(qn('w:type'),'dxa'); margins.append(e)
  cp.append(margins)
  if kind!='detail':
   sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'E8EEF5' if kind=='header' else 'F2F4F7'); cp.append(sh)
  for para in c.paragraphs:
   para.paragraph_format.space_before=Pt(0); para.paragraph_format.space_after=Pt(0); para.paragraph_format.line_spacing=1
   para.paragraph_format.keep_with_next=(kind in ['group','header'])
   for run in para.runs: run.font.size=Pt(10); run.bold=kind!='detail'
row(['STT','Thiết bị','Ghi chú'],'header')
n=0
for name,count in items:
 row(['',name,f'Số lượng: {count}'],'group')
 for _ in range(count):
  n+=1; row([str(n),name,''])
d.save(p)
c=Document(p)
assert len(c.tables)==1 and len(c.tables[0].rows)==102 and n==90
assert sum('Số lượng:' in r.cells[2].text for r in c.tables[0].rows)==11
print('Verified: 1 table, 11 group rows, 90 device rows.')
